"""Append government-facing pilot requirements to the existing team brief once."""

from pathlib import Path

from docx import Document


DOCUMENT_PATH = Path(
    r"C:\Users\Rupesh\Documents\Codex\2026-08-25\referenced-chatgpt-conversation-this-is-an"
) / "AirSentinel_Team_Workflow_and_Plan.docx"
HEADING = "Government-facing pilot requirements"


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def add_number(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Number")


document = Document(DOCUMENT_PATH)
if any(paragraph.text == HEADING for paragraph in document.paragraphs):
    print("Government-facing section already exists; no change made.")
else:
    document.add_page_break()
    document.add_heading(HEADING, level=1)
    document.add_paragraph(
        "AirSentinel should complement CPCB, SPCB/PCC, ULB, district and NCAP "
        "workflows. It is not a replacement for official AQI, statutory monitoring, "
        "enforcement authority or laboratory verification. The team must not make "
        "unverified claims about cybersecurity weaknesses in government systems and must "
        "never attempt to access, scan or bypass them."
    )

    document.add_heading("What AirSentinel adds", level=2)
    add_bullet(document, "Local evidence retained alongside city summaries: station, hotspot grid and locality confidence.")
    add_bullet(document, "An operational anomaly queue: case ID, evidence, confidence, suggested next step, assigned owner, due time and outcome.")
    add_bullet(document, "A clear separation between the public view and an authority view with role-based access.")
    add_bullet(document, "Auditable model/rule versions, data freshness and uncertainty rather than unexplained scores.")

    document.add_heading("Government operator case flow", level=2)
    add_number(document, "A verified measurement, anomaly or corroborated report creates a case with an evidence snapshot.")
    add_number(document, "The system displays source tier, data freshness, confidence and recommended review step.")
    add_number(document, "A human operator assigns the case, records the action and closes it with an outcome.")
    add_number(document, "The team evaluates forecast accuracy, anomaly false alerts, coverage and time to triage from completed cases.")

    document.add_heading("Safeguards the pilot must demonstrate", level=2)
    add_bullet(document, "A city average must not hide station/local hotspot variation; show location resolution and confidence.")
    add_bullet(document, "Sparse coverage must remain visible; never create an official AQI where evidence is missing.")
    add_bullet(document, "A forecast miss must create an unexpected-spike case, not be silently ignored.")
    add_bullet(document, "One anonymous report cannot trigger enforcement or public source attribution; use rate limits, moderation and corroboration.")
    add_bullet(document, "Indoor purifier data must be labelled separately from outdoor monitoring and accepted only with consent.")
    add_bullet(document, "No private address, account token or raw device ID may appear in analytics or the public dashboard.")
    add_bullet(document, "Automated suggestions require human review before enforcement or public attribution.")

    document.save(DOCUMENT_PATH)
    print(f"Updated: {DOCUMENT_PATH}")
