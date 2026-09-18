"""Build the submission-ready AirSentinel Word report.

The document is generated from verified repository facts. It deliberately
distinguishes local implementation from account-bound cloud deployment.
"""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"
OUTPUT = DOCS / "AirSentinel_Final_Project_Report.docx"

NAVY = "123047"
TEAL = "087E8B"
GREEN = "2E7D32"
LIGHT_TEAL = "E8F5F6"
LIGHT_GREEN = "EAF5EC"
LIGHT_BLUE = "EAF0F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
AMBER = "A15C00"
RED = "A12622"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("AIR SENTINEL  •  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 8),
        ("Heading 1", 20, NAVY, 16, 8),
        ("Heading 2", 14, TEAL, 12, 5),
        ("Heading 3", 11, GREEN, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Caption Small" not in styles:
        cap = styles.add_style("Caption Small", WD_STYLE_TYPE.PARAGRAPH)
        cap.font.name = "Aptos"
        cap.font.size = Pt(8.5)
        cap.font.italic = True
        cap.font.color.rgb = RGBColor.from_string(MID_GRAY)
        cap.paragraph_format.space_after = Pt(8)

    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.85)
        section.right_margin = Cm(1.85)


def add_label(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    r.font.name = "Aptos"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)


def add_callout(doc: Document, title: str, body: str, color=LIGHT_TEAL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    r.font.size = Pt(11)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(3)
        number = p.add_run(f"{index}.  ")
        number.bold = True
        number.font.color.rgb = RGBColor.from_string(TEAL)
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(WHITE)
        if widths:
            cell.width = Cm(widths[i])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            if row_index % 2:
                set_cell_shading(cell, LIGHT_GRAY)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string(NAVY)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/aptos-bold.ttf" if bold else "C:/Windows/Fonts/aptos.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def _centered(draw, box, label, font, fill="#123047"):
    lines = label.split("\n")
    spacing = 8
    sizes = [draw.textbbox((0, 0), line, font=font) for line in lines]
    heights = [b[3] - b[1] for b in sizes]
    total_h = sum(heights) + spacing * (len(lines) - 1)
    y = (box[1] + box[3] - total_h) / 2
    for line, bounds, height in zip(lines, sizes, heights):
        width = bounds[2] - bounds[0]
        x = (box[0] + box[2] - width) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += height + spacing


def _arrow(draw, start, end, color="#087E8B", width=6):
    draw.line([start, end], fill=color, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 18, y - 11), (x - 18, y + 11)], fill=color)


def draw_flow(path: Path) -> None:
    image = Image.new("RGB", (2400, 880), "white")
    draw = ImageDraw.Draw(image)
    colors = ["#EAF0F8", "#E8F5F6", "#EAF5EC", "#FFF2D6", "#FDECEC"]
    labels = ["Trusted\nevidence", "Quality &\ncoverage gates", "1h / 3h / 4h\nforecast", "Unexpected\nchange case", "Human review &\naudited action"]
    boxes = []
    x = 70
    for index, label in enumerate(labels):
        width = 390 if index < 4 else 420
        box = (x, 250, x + width, 500)
        boxes.append(box)
        draw.rounded_rectangle(box, radius=28, fill=colors[index], outline="#087E8B", width=5)
        _centered(draw, box, label, _font(46, True))
        if index:
            previous = boxes[index - 1]
            _arrow(draw, (previous[2] + 14, 375), (box[0] - 14, 375))
        x += width + 90
    title = "CAUSE UNVERIFIED  •  UNCERTAINTY PRESERVED  •  NO AUTOMATIC ENFORCEMENT"
    title_box = draw.textbbox((0, 0), title, font=_font(35, True))
    draw.text(((2400 - (title_box[2] - title_box[0])) / 2, 610), title, font=_font(35, True), fill="#A12622")
    note = "CPCB current data when authorised  |  OpenAQ historical fallback visibly labelled  |  Context is not proof"
    note_box = draw.textbbox((0, 0), note, font=_font(30))
    draw.text(((2400 - (note_box[2] - note_box[0])) / 2, 700), note, font=_font(30), fill="#667085")
    image.save(path)


def draw_cloud(path: Path) -> None:
    image = Image.new("RGB", (2400, 1050), "white")
    draw = ImageDraw.Draw(image)
    title = "PUBLIC AGGREGATES SEPARATED FROM PRIVATE EVIDENCE AND REVIEW ACTIONS"
    bounds = draw.textbbox((0, 0), title, font=_font(38, True))
    draw.text(((2400 - (bounds[2] - bounds[0])) / 2, 55), title, font=_font(38, True), fill="#123047")
    left_nodes = [
        ((80, 235, 510, 390), "Public aggregate\ndashboard", "#EAF0F8"),
        ((80, 465, 510, 620), "Citizen\ninterface", "#E8F5F6"),
        ((80, 695, 510, 850), "Authority\nportal", "#EAF5EC"),
    ]
    center = (785, 410, 1515, 675)
    right_nodes = [
        ((1780, 185, 2320, 335), "Firebase\nidentity + reports", "#E8F5F6"),
        ((1780, 385, 2320, 535), "BigQuery\ngoverned analytics", "#EAF0F8"),
        ((1780, 585, 2320, 735), "Gemini + language\nguarded helpers", "#FFF2D6"),
        ((1780, 785, 2320, 935), "Maps + Earth Engine\ncontext only", "#EAF5EC"),
    ]
    for box, label, color in left_nodes + right_nodes:
        draw.rounded_rectangle(box, radius=26, fill=color, outline="#087E8B", width=5)
        _centered(draw, box, label, _font(39, True))
    draw.rounded_rectangle(center, radius=30, fill="#DCEFF2", outline="#087E8B", width=6)
    _centered(draw, center, "Private Cloud Run API\nIAM + human-review guard", _font(48, True))
    for box, _, _ in left_nodes:
        _arrow(draw, (box[2] + 18, (box[1] + box[3]) // 2), (center[0] - 18, (center[1] + center[3]) // 2))
    for box, _, _ in right_nodes:
        _arrow(draw, (center[2] + 18, (center[1] + center[3]) // 2), (box[0] - 18, (box[1] + box[3]) // 2))
    image.save(path)


def add_cover(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()
    add_label(doc, "BRICS Track 2 • Sustainability • India-first prototype")
    p = doc.add_paragraph(style="Title")
    p.add_run("AirSentinel")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(20)
    r = p2.add_run("Locality-aware clean-air early warning\nand human authority review")
    r.font.name = "Aptos Display"
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    r.font.bold = True

    add_callout(
        doc,
        "Working prototype, evidence-first claims",
        "Six-pollutant monitoring • 181.9-day locality evidence • 1h/3h/4h forecasts • sudden-spike review • citizen multimodal support • secure Google Cloud path",
        LIGHT_GREEN,
    )
    doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(60)
    p3.add_run("FINAL PROJECT REPORT").bold = True
    p4 = doc.add_paragraph("26 August 2026")
    p4.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    p5 = doc.add_paragraph("Status: local implementation verified; cloud deployment requires team-owned credentials and billing.")
    p5.runs[0].font.italic = True
    p5.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    doc.add_page_break()


def build() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    flow_path = ASSETS / "final_evidence_to_action.png"
    cloud_path = ASSETS / "final_google_cloud_architecture.png"
    draw_flow(flow_path)
    draw_cloud(cloud_path)

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.header.paragraphs[0].text = "AIR SENTINEL  |  CLEAN AIR & CLIMATE RESILIENCE"
    section.header.paragraphs[0].runs[0].font.size = Pt(8)
    section.header.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_page_number(section.footer.paragraphs[0])

    add_cover(doc)

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "AirSentinel closes the gap between city-wide air-quality visibility and a local government decision. "
        "It combines provenance-labelled outdoor observations, weather, contextual satellite layers and consented "
        "community evidence; forecasts one to four hours ahead; detects unsupported coverage or a sudden forecast "
        "miss; and creates a cause-unverified case for authenticated human review."
    )
    add_callout(
        doc,
        "Main result",
        "Najafgarh retains 7,516 PM2.5 rows and Safdarjung 7,889 rows across 181.9 days, with two reporting stations each. Their stored PM2.5 completeness is 86.1% and 90.4%, respectively. The source is labelled OpenAQ fallback data, not an authorised CPCB archive.",
    )
    doc.add_heading("What judges can evaluate end to end", level=2)
    add_numbered(doc, [
        "Inspect national/city conditions and locality evidence quality.",
        "View a one-, three- or four-hour historical prototype forecast and uncertainty range.",
        "Open an unexpected-spike or evidence-gap case whose cause remains unverified.",
        "Submit privacy-minimised citizen evidence and keep it pending moderation.",
        "Record a local demonstration review action, or use the private authorised API path for production-mode action testing.",
        "Inspect honest Google integration readiness and secure deployment boundaries.",
    ])
    doc.add_heading("Truth boundary", level=2)
    add_bullets(doc, [
        "Not official AQI, a medical device, source attribution or automatic enforcement.",
        "OpenAQ is a labelled historical-development fallback, not presented as an authorised CPCB archive.",
        "No Google Cloud deployment is claimed until a team-owned project and verified URLs exist.",
        "Deep-learning and federated-learning models are not promoted in this submission.",
    ])

    doc.add_heading("1. Problem and solution", level=1)
    doc.add_heading("The four gaps", level=2)
    add_table(doc, ["Gap", "Why it matters", "AirSentinel response"], [
        ["Scale", "A city value can hide local events.", "Separate locality anchors/geometries and evidence gates."],
        ["Coverage", "Sparse or stale stations create false precision.", "Show evidence gaps, station provenance, distance and completeness."],
        ["Forecast failure", "A sudden event may lie outside the expected range.", "Create a cause-unverified review case and seek corroboration."],
        ["Coordination", "Station, weather, satellite and citizen data use different trust levels.", "Canonical records, contextual labels, identity and audit."],
    ], [2.4, 6.0, 8.2])
    doc.add_picture(str(flow_path), width=Inches(6.9))
    p = doc.add_paragraph("Figure 1. Evidence is gated before forecasting; a sudden mismatch ends in human review, not automated blame.", style="Caption Small")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("2. Data improvement and trust", level=1)
    doc.add_paragraph(
        "The historical collector now supports date chunks, pagination, retries, active-sensor probes, two working sensors per locality/pollutant, deduplication and safe merging. It preserves source-system labels so a coverage gate cannot silently upgrade a fallback into an official source."
    )
    add_table(doc, ["Locality", "PM2.5 rows", "Span", "Completeness", "Stations", "Pollutants"], [
        ["Najafgarh pilot area", "7,516", "181.9 days", "86.1%", "2", "PM2.5, PM10, NO2, O3, SO2, CO"],
        ["Safdarjung Enclave pilot area", "7,889", "181.9 days", "90.4%", "2", "PM2.5, PM10, NO2, O3, SO2, CO"],
    ], [3.8, 1.9, 1.9, 2.5, 1.7, 5.0])
    add_callout(
        doc,
        "Evidence gate passed—not operational promotion",
        "Both localities now support a controlled conventional baseline comparison. They still need approved geometry, station mapping, calibration/uptime context, weather alignment and fresh chronological evaluation before operational use.",
        LIGHT_BLUE,
    )
    doc.add_heading("Source hierarchy and alternatives", level=2)
    add_table(doc, ["Priority", "Source", "Use", "Limit"], [
        ["1", "CPCB/data.gov.in current catalogue", "Preferred official Indian current observations", "Authorised key/configuration required; current data is not a guaranteed historical archive."],
        ["2", "CPCB/SPCB/PCC approved export", "Historical production evidence", "Requires authority access, calibration and station metadata."],
        ["Fallback", "OpenAQ", "Historical development and discovery", "Keep provider provenance and visible fallback label."],
        ["Context", "Open-Meteo / Earth Engine / FIRMS", "Weather, satellite gas/aerosol and fire context", "Cannot independently prove surface PM2.5 or a source."],
    ], [2.0, 4.0, 4.8, 5.9])

    doc.add_heading("3. Forecasting and sudden-event handling", level=1)
    add_bullets(doc, [
        "Direct one-, three- and four-hour targets rather than recursively compounding a one-hour model.",
        "Leakage-safe lag, rolling, weather and available pollutant features.",
        "Time-ordered training/validation and newest untouched hold-out.",
        "Persistence first; established tree boosting only when it demonstrates stable improvement.",
        "Empirical residual ranges and an explicit research/prototype label.",
    ])
    doc.add_heading("What happens when the model is wrong", level=2)
    add_numbered(doc, [
        "Detect that a fresh observation lies well outside the expected residual distribution.",
        "Create an unexpected local spike case and preserve the observation/forecast evidence.",
        "Mark the cause unverified; do not infer burning, industry, traffic or sensor failure.",
        "Ask a reviewer to check nearby stations, sensor health, weather, satellite/fire context and moderated citizen signals.",
        "Record the reviewer’s monitor/inspect/escalate/close decision in the audit trail.",
    ])
    add_callout(doc, "Model policy", "The submission does not promote deep learning or federated learning. Vertex AI is only an optional host for a conventional, human-approved model that has earned promotion on approved data.", "FFF2D6")

    doc.add_heading("4. Working system components", level=1)
    add_table(doc, ["Component", "Job", "Safety boundary"], [
        ["Public dashboard", "Reviewed city/locality aggregates, coverage and forecast ranges", "No personal evidence or enforcement actions"],
        ["Citizen interface", "Consented coarse-area text/photo/voice reports", "Unverified pending evidence; exact location minimised"],
        ["Authority portal", "Case evidence and reviewer ownership; production UI is view-only", "Writes use private authorised API; no automated external action"],
        ["FastAPI", "Typed contracts, readiness and integration endpoints", "Fail-closed production authentication"],
        ["Data pipeline", "Discovery, collection, validation, provenance and coverage", "Fallback labels cannot be promoted by volume"],
        ["Forecast/anomaly", "Short-horizon estimates and residual review cases", "Historical prototype and cause-unverified language"],
    ], [3.3, 7.0, 6.5])

    doc.add_heading("5. Google technology architecture", level=1)
    doc.add_picture(str(cloud_path), width=Inches(6.9))
    p = doc.add_paragraph("Figure 2. Public aggregates are separated from private evidence, identity and authority actions.", style="Caption Small")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_table(doc, ["Technology", "Meaningful job", "Activation requirement"], [
        ["Gemini API / AI Studio", "Guarded case explanation and in-memory photo description", "Secret Manager value, feature review and output audit"],
        ["Vertex AI", "Optional conventional-model hosting/monitoring", "Approved model, model card and monitoring plan"],
        ["Firebase", "Citizen identity and minimum pending-review metadata", "Project, Auth, Firestore, rules, retention and abuse controls"],
        ["BigQuery", "Governed coverage/readiness/reviewed analytical tables", "Dataset region, IAM, retention and cost policy"],
        ["Maps / Earth Engine", "Pilot-location and satellite contextual layers", "Restricted key / registered project; contextual-only policy"],
        ["STT / Translation / TTS", "Accessible multilingual report and guidance workflow", "Consent, language QA, retention and enabled APIs"],
        ["Cloud Run", "Public aggregate, private API and private authority services", "Billing, IAM, secrets and URL/identity tests"],
    ], [3.6, 7.3, 6.3])

    doc.add_heading("6. Privacy, security and human governance", level=1)
    add_bullets(doc, [
        "Verified citizen identity in production; stronger reviewer identity for authority operations.",
        "Deny-by-default direct Firestore and Storage rules; server access through least-privilege cloud identities.",
        "No exact private address, raw media or unnecessary contact data in analytical tables.",
        "Bounded image/audio types and sizes; Gemini output cannot claim AQI, liability, enforcement or health diagnosis.",
        "Secrets stay in Secret Manager and are excluded from Git, Docker context and documentation.",
        "Production authority dashboard actions are view-only; writes must pass the private API's authenticated role policy and require durable storage before a pilot.",
        "The public Cloud Run image is built from a temporary six-file aggregate allowlist; private queues, moderation summaries and model features are excluded.",
        "Thirty-eight offline tests pass, including real in-process ASGI request-boundary checks and fail-closed cloud-adapter tests.",
    ])
    add_callout(doc, "Government-facing operating rule", "AirSentinel may recommend what to verify next. An authorised officer decides whether to monitor, inspect, coordinate, issue guidance or close the case through the appropriate official process.", LIGHT_GREEN)

    doc.add_heading("7. Revised phases and ownership", level=1)
    add_table(doc, ["Phase", "Delivered", "External owner action"], [
        ["0 Foundation", "Project, exclusions, evidence language and tests", "Nominate data/privacy owner"],
        ["1 Trusted data", "Six-pollutant sources, contract and QA", "Authorised CPCB history/current access"],
        ["2 Long history", "Najafgarh/Safdarjung 181.9-day fallback evidence", "Verify mapping/calibration/geometry"],
        ["3 Forecast & anomaly", "1h/3h/4h research candidates and residual cases", "Fresh approved-data re-evaluation"],
        ["4 Authority review", "Private queue, actions and audit", "Official roles and escalation SLA"],
        ["5 Citizen/accessibility", "Text/photo/voice and moderation contracts", "Firebase Auth, consent, retention, abuse response"],
        ["6 Google services", "Adapters and readiness checks", "Project, APIs, secrets and quotas"],
        ["7 Deployment", "Docker and Cloud Run scripts", "Run/test real deployment and identities"],
        ["8 Submission", "Reports, deck, script and checklist", "Push GitHub, record video, insert URLs"],
    ], [3.0, 7.0, 7.0])

    doc.add_heading("8. Deployment and operational path", level=1)
    add_numbered(doc, [
        "Create/select a team-owned GCP project, billing account, region and retention policy.",
        "Authenticate gcloud and run the least-privilege bootstrap script.",
        "Add reviewed secrets directly to Secret Manager and configure Firebase Auth/rules.",
        "Create BigQuery governed tables and upload only approved analytical artifacts.",
        "Deploy the private API and private authority portal from the private image; deploy the public dashboard from its temporary allowlisted context.",
        "Test citizen identity, reviewer identity, audit persistence, negative access cases and all generated URLs.",
        "Replace fallback/local file paths with approved persistent services before any operational pilot.",
    ])
    add_callout(doc, "Current deployment truth", "No GCP project, billing account, gcloud authentication or public URL exists in the local environment. The code and scripts are ready; the account-bound execution remains a team owner task.", "FDECEC")

    doc.add_heading("9. Demo and submission package", level=1)
    add_table(doc, ["Required item", "Prepared artifact", "Final human action"], [
        ["Source repository", "Secret-safe source tree and release manifest", "Create/push GitHub and add URL"],
        ["3–5 minute video", "Timed script and demo runbook", "Record narration/screen and upload"],
        ["10–12 slide deck", "AirSentinel_Hackathon_Pitch_Deck.pptx", "Add final URLs/team names if needed"],
        ["Short description", "SHORT_PROJECT_DESCRIPTION.md", "Paste into submission form"],
        ["Deployed link", "Cloud Run scripts and status template", "Deploy/test and add verified URL"],
    ], [3.8, 7.0, 6.2])
    doc.add_heading("Recommended demo story", level=2)
    add_numbered(doc, [
        "Show the local monitoring gap and evidence quality.",
        "Show Najafgarh/Safdarjung’s improved historical coverage.",
        "Show a short-horizon forecast and uncertainty.",
        "Show a forecast miss becoming a cause-unverified case.",
        "Show citizen evidence and a human authority action.",
        "Show the Google service boundary and name the remaining account-bound steps honestly.",
    ])

    doc.add_heading("10. Limitations and next decisions", level=1)
    add_bullets(doc, [
        "Fallback history is not an approved government archive.",
        "Locality pilots are station-proxy areas until geometry and physical mapping are approved.",
        "Historical forecast performance is not live operational performance.",
        "Satellite vertical-column products are not surface PM2.5 in micrograms per cubic metre.",
        "Citizen, purifier and photo evidence supports triage but does not substitute for verified outdoor monitoring.",
        "Cloud Run local disk is ephemeral; operational reports/actions require Firestore/BigQuery persistence.",
        "Public warnings, health advice, source attribution and enforcement need the competent authority’s policy and approval.",
    ])

    doc.add_heading("11. Source references", level=1)
    sources = [
        "CPCB/data.gov.in real-time air-quality catalogue — https://www.data.gov.in/catalog/real-time-air-quality-index",
        "OpenAQ API documentation — https://docs.openaq.org/",
        "Open-Meteo historical weather — https://open-meteo.com/en/docs/historical-weather-api",
        "Google Cloud Run authentication — https://cloud.google.com/run/docs/authenticating/overview",
        "Firebase security rules and authentication — https://firebase.google.com/docs/rules/rules-and-auth",
        "Gemini image understanding — https://ai.google.dev/gemini-api/docs/image-understanding",
        "Earth Engine Sentinel-5P catalogue — https://developers.google.com/earth-engine/datasets/catalog/sentinel-5p",
        "NASA FIRMS — https://firms.modaps.eosdis.nasa.gov/",
    ]
    add_bullets(doc, sources)

    doc.add_heading("Appendix A. Local commands", level=1)
    commands = [
        r"cd C:\Users\Rupesh\AirSentinel\AirSentinel",
        r".\.venv\Scripts\python.exe -m unittest discover -s tests -v",
        r".\.venv\Scripts\python.exe -m streamlit run app.py",
        r".\.venv\Scripts\python.exe -m streamlit run authority_app.py --server.port 8502",
        r".\.venv\Scripts\python.exe -m streamlit run citizen_app.py --server.port 8503",
        r".\.venv\Scripts\python.exe -m uvicorn src.api.main:app --port 8000",
    ]
    for command in commands:
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_margins(cell)
        run = cell.paragraphs[0].add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    doc.add_heading("Appendix B. Submission message", level=1)
    add_callout(
        doc,
        "AirSentinel",
        "A locality-aware clean-air early-warning and human-review platform that combines provenance-labelled station data, weather, satellite context and consented community evidence to detect unsupported coverage and unexpected pollution changes, forecast one to four hours ahead, and guide an auditable sustainable response—without inventing hyper-local precision or automating blame.",
        LIGHT_TEAL,
    )

    core = doc.core_properties
    core.title = "AirSentinel Final Project Report"
    core.subject = "BRICS Track 2 — Clean Air & Climate Resilience"
    core.author = "AirSentinel Team"
    core.keywords = "air quality, sustainability, India, locality, Google Cloud, human review"
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
