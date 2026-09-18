"""Create three beginner-friendly AirSentinel Word guides from verified project facts."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"

NAVY = "123047"
BLUE = "2E74B5"
TEAL = "087E8B"
GREEN = "2E7D32"
INK = "1F2937"
MUTED = "667085"
PALE_BLUE = "E8EEF5"
PALE_TEAL = "E8F5F6"
PALE_AMBER = "FFF4D6"
PALE_GREEN = "EAF5EC"
LIGHT_GREY = "F4F6F9"
WHITE = "FFFFFF"

OUT_TEAM = DOCS / "AirSentinel_Remaining_Team_Work_Guide.docx"
OUT_LEARN = DOCS / "AirSentinel_Code_Learning_Handbook.docx"
OUT_SYSTEM = DOCS / "AirSentinel_System_Working_Explained.docx"


def shade(cell, colour):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        pr.append(node)
    node.set(qn("w:fill"), colour)


def margins(cell, top=80, start=120, bottom=80, end=120):
    pr = cell._tc.get_or_add_tcPr()
    tc_mar = pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_split(row):
    pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    pr.append(node)


def repeating_header(row):
    pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    pr.append(node)


def fixed_table(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("AirSentinel  •  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def setup(doc, subtitle):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(8)

    for name, size, colour, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(colour)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.text = subtitle.upper()
    header.runs[0].font.name = "Calibri"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)


def cover(doc, title, strapline, purpose, contents):
    doc.add_paragraph("AIR SENTINEL", style="Heading 3")
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(strapline)
    p.runs[0].font.size = Pt(15)
    p.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    p.paragraph_format.space_after = Pt(18)
    callout(doc, "Purpose", purpose, PALE_TEAL)
    doc.add_heading("What is inside", 2)
    bullets(doc, contents)
    p = doc.add_paragraph("Prepared from the verified local AirSentinel repository. It clearly separates what already works from work that needs your team’s accounts, approvals or field validation.")
    p.paragraph_format.space_before = Pt(18)
    p.runs[0].italic = True
    doc.add_page_break()


def callout(doc, title, text, colour=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    fixed_table(table, [6.5])
    cell = table.cell(0, 0)
    shade(cell, colour)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run("\n" + text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def code_block(doc, command):
    t = doc.add_table(rows=1, cols=1)
    fixed_table(t, [6.5])
    cell = t.cell(0, 0)
    shade(cell, LIGHT_GREY)
    r = cell.paragraphs[0].add_run(command)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, headers):
        shade(cell, PALE_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].add_run(text).bold = True
    repeating_header(t.rows[0])
    for values in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        prevent_split(t.rows[-1])
    fixed_table(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def simple_flow(doc, title, steps):
    doc.add_heading(title, 2)
    for index, (name, text) in enumerate(steps, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        badge = p.add_run(f"{index}. ")
        badge.bold = True
        badge.font.color.rgb = RGBColor.from_string(TEAL)
        name_run = p.add_run(name + " — ")
        name_run.bold = True
        p.add_run(text)


def add_image(doc, path, caption, width=6.25):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def team_guide():
    doc = Document()
    setup(doc, "Remaining team work guide")
    cover(doc, "Remaining Team Work Guide", "A step-by-step plan to finish AirSentinel together", "This guide tells the team exactly what still needs people, accounts, approvals and deployment. Start at Step 1 and do not claim an external integration is live until its checkpoint passes.", ["A clear ownership split", "Data improvement work", "Google Cloud set-up and deployment", "GitHub, video and submission steps", "Definitions of done and troubleshooting"])

    doc.add_heading("1. What is already complete", 1)
    doc.add_paragraph("The local prototype works: public, citizen and authority dashboards; a private FastAPI backend; six-pollutant data contracts; locality evidence gates; one-, three- and four-hour historical forecast candidates; anomaly review; audit records; feature-flagged Google adapters; deployment scripts; and 38 passing tests.")
    callout(doc, "Important truth", "The project is not yet a live government system. OpenAQ is a labelled development fallback, current CPCB/data.gov.in access is not configured, and Google Cloud cannot be deployed until your team uses its own project, billing and identities.", PALE_AMBER)

    doc.add_heading("2. Give each person one clear owner role", 1)
    table(doc, ["Role", "Owns", "Finish line"], [
        ["Project/cloud owner", "GCP project, billing, region, IAM, Secret Manager, Cloud Run", "Public/private URLs work and access is least privilege."],
        ["Data owner", "CPCB/data.gov.in access, official archive requests, station mapping, units and calibration", "Every data source has provenance and locality relationship evidence."],
        ["ML owner", "Coverage gate, chronological baselines, uncertainty and drift checks", "Najafgarh/Safdarjung are evaluated without lowering safety thresholds."],
        ["Citizen/governance owner", "Consent, Firebase Auth, moderation, retention and reviewer roles", "Reports cannot be mistaken for official measurements."],
        ["Demo/release owner", "GitHub, deployed links, screenshots, pitch, 3–5 minute video", "Judges can evaluate the whole flow end to end."],
    ], [1.4, 3.1, 2.0])

    doc.add_heading("3. Work in this order", 1)
    simple_flow(doc, "Do not skip this sequence", [
        ("Create ownership", "Name the five owners above and put the shared project/contact details in a team document."),
        ("Secure data", "Request and document the approved current/historical source before treating the model as operational."),
        ("Configure cloud", "Create one team-owned GCP project, enable billing and add restricted service identities."),
        ("Deploy safely", "Deploy public dashboard separately from private API and authority dashboard."),
        ("Test identities", "Test allowed and denied actions, persistence and audit logs."),
        ("Package proof", "Push GitHub, record video and submit only verified URLs and claims."),
    ])

    doc.add_heading("4. Data improvement: the highest-priority team task", 1)
    doc.add_paragraph("The system requires six outdoor pollutants: PM2.5, PM10, NO₂, O₃, SO₂ and CO. The preferred production path is an authorised Indian source such as CPCB/data.gov.in or an approved CPCB/SPCB/PCC export. The current long history is a clearly marked OpenAQ development fallback; it is useful for research, not proof of official operational readiness.")
    table(doc, ["Locality", "Available fallback history", "What still needs validation"], [
        ["Najafgarh pilot area", "7,516 PM2.5 rows; 181.9 days; 86.1% completeness; two stations", "Approved boundary, station-to-locality mapping, calibration/uptime and a controlled baseline evaluation."],
        ["Safdarjung Enclave pilot area", "7,889 PM2.5 rows; 181.9 days; 90.4% completeness; two stations", "Approved boundary, station-to-locality mapping, calibration/uptime and a controlled baseline evaluation."],
    ], [1.55, 2.55, 2.4])
    numbered(doc, [
        "Ask CPCB/SPCB/PCC or the authorised data.gov.in contact for current observations, a historical export, station metadata, units, calibration/maintenance and uptime information.",
        "Create a station mapping sheet: station ID, coordinates, source, operating status, distance to locality, whether it is inside the approved locality boundary, and responsible authority.",
        "Keep the phrase ‘nearby-station proxy’ until an authority approves the locality geometry and physical station relationship.",
        "Store approved analytical tables only. Never commit raw downloads, personal reports, credentials or audit data to GitHub.",
        "Re-run the repository coverage and promotion checks after new data arrives. Do not lower the 240 safe-row threshold simply to get a result.",
    ])
    code_block(doc, r"cd C:\Users\Rupesh\AirSentinel\AirSentinel")
    code_block(doc, r".\.venv\Scripts\python.exe src\data\report_multiscale_coverage.py")
    code_block(doc, r".\.venv\Scripts\python.exe src\models\report_model_promotion_readiness.py")
    callout(doc, "Current model gate", "Najafgarh has 224/207/202 safe aligned evaluation rows for 1h/3h/4h; Safdarjung has 151/135/131. The controlled requirement is 240 per horizon, so both remain historical prototype candidates rather than promoted local models.", PALE_AMBER)

    doc.add_heading("5. Google Cloud work: complete this as one team-owned setup", 1)
    doc.add_paragraph("Use one Google Cloud project owned by the team, not a personal demo project. Choose the region, owner, data retention policy and billing owner before deploying. Keep secrets in Secret Manager, never in a source file or screenshot.")
    table(doc, ["Service", "Why AirSentinel uses it", "Team action"], [
        ["Cloud Run", "Hosts public dashboard separately from private services.", "Create project/billing; deploy and test each URL."],
        ["Firebase", "Verified citizen identity and pending-review evidence.", "Enable Auth, apply deny-by-default rules, set consent/retention."],
        ["BigQuery", "Approved analytics/audit reporting at scale.", "Create dataset and least-privilege writer/viewer roles."],
        ["Gemini", "Bounded explanation and photo evidence triage.", "Enable API, set key/identity and keep human review."],
        ["Maps / Earth Engine", "Map context and satellite context—not source proof.", "Restrict Maps key; register Earth Engine project."],
        ["Translation / STT / TTS", "Accessibility for consented citizen reports.", "Enable only if required, test target languages and consent text."],
    ], [1.15, 2.75, 2.6])
    doc.add_heading("Safe deployment checklist", 2)
    numbered(doc, [
        "Read deploy/README_GOOGLE_CLOUD.md together before running a command.",
        "Create/select the team GCP project. Enable billing and record its ID in the team password manager/document—not in public code.",
        "Install and authenticate gcloud on the cloud owner’s computer. Use separate least-privilege identities for deployment and runtime.",
        "Run the repository bootstrap scripts with the real project ID. Review every resource before allowing it to be created.",
        "Add reviewed secret values directly in Google Secret Manager. Do not use .env values in Cloud Run command arguments.",
        "Deploy three separated surfaces: public aggregate dashboard, private API, and private authority dashboard.",
        "Test one citizen identity, one reviewer identity, one denied/unprivileged user, audit persistence and service restart behaviour.",
    ])
    code_block(doc, r"Get-Content deploy\README_GOOGLE_CLOUD.md")
    code_block(doc, r"powershell -ExecutionPolicy Bypass -File deploy\gcp-bootstrap.ps1 -ProjectId <YOUR_PROJECT_ID>")
    code_block(doc, r"powershell -ExecutionPolicy Bypass -File deploy\cloud-run-deploy.ps1 -ProjectId <YOUR_PROJECT_ID> -Region <YOUR_REGION>")
    callout(doc, "Do not invent a production action", "Cloud Functions, Dialogflow and any external notification are optional until the hackathon rules and an approved event contract give you an owner, consent policy and idempotent trigger. The existing system correctly avoids pretending these are live.", PALE_AMBER)

    doc.add_heading("6. Deployment acceptance tests", 1)
    table(doc, ["Test", "Expected result", "Pass condition"], [
        ["Public dashboard", "Only aggregate, allowlisted information is shown.", "No case action, raw report or contact data is exposed."],
        ["Citizen report", "Verified user can submit coarse, consented evidence.", "Submission enters pending review; it is not treated as a sensor measurement."],
        ["Authority action", "Verified reviewer can view and close a case.", "The action is durable and audit-log fields are stored."],
        ["Negative access", "Unauthorised user is refused.", "Private endpoints and Firestore/Storage client access deny by default."],
        ["Failure handling", "A disabled Google adapter returns a clear safe message.", "No fake Gemini/Maps/Earth Engine result is shown."],
    ], [1.25, 3.0, 2.25])

    doc.add_heading("7. GitHub and submission package", 1)
    numbered(doc, [
        "Create a new GitHub repository under a team-owned account or organisation.",
        "Run the tests and inspect the status before the first commit. Ensure .env, service-account files, raw data, citizen reports, audit records and model binaries are excluded.",
        "Push only the curated source tree and approved documentation. The repository currently has a curated submission ZIP but GitHub publication and remote setup are team work.",
        "Record the 3–5 minute video using docs/DEMO_VIDEO_SCRIPT.md and docs/DEMO_RUNBOOK.md. Show a real end-to-end flow, its safety boundaries and the final deployed link.",
        "Open docs/SUBMISSION_CHECKLIST.md and tick every external link, team name and required form field only after verification.",
    ])
    code_block(doc, r".\.venv\Scripts\python.exe -m unittest discover -s tests -v")
    code_block(doc, r"git status")
    code_block(doc, r"git add .")
    code_block(doc, r"git commit -m \"Prepare AirSentinel hackathon submission\"")
    callout(doc, "Before pressing submit", "The source URL, demo URL, deployed public URL, team details and 3–5 minute video must be real and tested. Do not say ‘deployed’, ‘official’, ‘live alert’ or ‘source identified’ unless your evidence supports it.", PALE_GREEN)

    doc.add_heading("8. A two-day finishing plan", 1)
    table(doc, ["When", "Owner", "Deliverable"], [
        ["Day 1 morning", "Cloud + data", "GCP project/billing/roles; official-source request; station mapping sheet."],
        ["Day 1 afternoon", "Cloud + governance", "Firebase/BigQuery/secrets configured; consent, roles and retention reviewed."],
        ["Day 2 morning", "Cloud + QA", "Cloud Run deployments; positive/negative identity tests; audit persistence test."],
        ["Day 2 afternoon", "Release + demo", "GitHub push, video recording, pitch/deployed links and submission form review."],
    ], [1.4, 1.7, 3.4])

    doc.add_heading("9. When you are genuinely done", 1)
    bullets(doc, [
        "The deployed public view is reachable and excludes sensitive/authority-only material.",
        "The authority view and API reject non-reviewer users, and every reviewed action is audited durably.",
        "Official data status, fallback status, freshness, locality relationship and uncertainty are visible.",
        "Your team can explain that a sudden spike becomes a reviewed case, not an automatic accusation or enforcement event.",
        "GitHub, demo video, pitch deck, short description and deployed link are all included and checked.",
    ])
    doc.core_properties.title = "AirSentinel Remaining Team Work Guide"
    doc.core_properties.author = "AirSentinel Team"
    doc.save(OUT_TEAM)


def learning_guide():
    doc = Document()
    setup(doc, "Code learning handbook")
    cover(doc, "Code Learning Handbook", "Understand the AirSentinel code well enough to explain and extend it", "This handbook is a beginner’s map from Python files to the working product. It focuses on how the real repository works, not on memorising every line.", ["A 2–4 hour preparation route", "What every folder does", "How data becomes a forecast and a case", "Security, Google adapters and testing", "Practical questions judges may ask"])

    doc.add_heading("1. Learn the system in the right order", 1)
    table(doc, ["Time", "Study", "Result"], [
        ["0–30 minutes", "README.md, HANDOFF.md and the public dashboard", "You can state the goal, non-claims and user roles."],
        ["30–75 minutes", "Data scripts and processed CSV outputs", "You know where observations come from and why provenance matters."],
        ["75–120 minutes", "Features, model training and forecast output", "You can explain lags, chronological tests, MAE/RMSE and uncertainty."],
        ["120–165 minutes", "Anomaly/case pipeline and authority app", "You can explain human review and audit safety."],
        ["165–240 minutes", "Cloud adapters, deployment guide and tests", "You can distinguish local working code from account-bound cloud activation."],
    ], [1.0, 2.75, 2.75])
    callout(doc, "Best learning method", "Run a small command, open the file that produced the output, and answer three questions: what is the input, what transformation happens, and what output is saved? This is better than trying to read the entire codebase from top to bottom.", PALE_TEAL)

    doc.add_heading("2. The project map", 1)
    table(doc, ["Path", "Plain-English purpose"], [
        ["app.py", "Public aggregate dashboard. It shows reviewed, non-sensitive information."],
        ["citizen_app.py", "Consent-based citizen text/photo/voice reporting interface."],
        ["authority_app.py", "Reviewer dashboard. Local prototype actions are audited; production mode is view-only."],
        ["src/api/", "FastAPI backend, request schemas and access checks for production-like protected actions."],
        ["src/data/", "Downloads, normalises, validates and reports station/pollutant data."],
        ["src/features/", "Builds time-series features such as lag values and weather variables."],
        ["src/models/", "Trains, compares, predicts and documents conventional locality models."],
        ["src/scoring/", "Turns an unexpected forecast mismatch into a review case—not a source claim."],
        ["src/authority/ and src/services/", "Case queue, audit record, citizen report and guarded Gemini helpers."],
        ["src/cloud/", "Feature-flagged adapters for Firebase, BigQuery, Gemini, Maps, Earth Engine and language services."],
        ["deploy/", "Cloud Run/GCP bootstrap and safe packaging scripts."],
        ["tests/", "Automated checks proving safety, source labels, access and data gates."],
    ], [1.8, 4.7])

    doc.add_page_break()
    doc.add_heading("3. Python ideas you are using", 1)
    table(doc, ["Idea", "Where it appears", "Why it matters"], [
        ["Modules and imports", "Every src/ folder", "Lets the app reuse one reliable function instead of duplicating logic."],
        ["Pathlib Path", "Data and model scripts", "Builds safe Windows/Linux file paths such as data/processed/output.csv."],
        ["Functions", "fetch_*, build_*, train_*", "A named, testable unit: input → work → output."],
        ["Pandas DataFrame", "Data/feature/model scripts", "A spreadsheet-like table that can filter, join, group and save CSV."],
        ["Requests", "Data connectors", "Calls an API and receives JSON, with timeout/error handling."],
        ["Environment variables", ".env / cloud adapters", "Keep secrets outside source code; feature flags make optional services safe."],
        ["Pydantic schemas", "src/api/schemas.py", "Check that API requests have the expected fields and types."],
        ["Tests", "tests/", "Verify assumptions repeatedly instead of trusting a manual demo."],
    ], [1.3, 2.1, 3.1])

    doc.add_heading("4. How data becomes an AirSentinel output", 1)
    simple_flow(doc, "The core pipeline", [
        ("Discover", "Find nearby stations for cities/locality pilot anchors and the six target pollutants."),
        ("Fetch", "Download latest/historical observations and weather with pagination, retries and provenance."),
        ("Validate", "Reject malformed records, standardise units/time, deduplicate and quantify coverage gaps."),
        ("Build features", "Make lag/rolling/time/weather features from prior observations only."),
        ("Evaluate", "Train a conventional model on earlier time; test it on later time to avoid leakage."),
        ("Predict", "Create 1h, 3h and 4h historical prototype forecast candidates with empirical intervals."),
        ("Review", "If actual data later disagrees strongly, create a cause-unverified case for a person to review."),
    ])
    code_block(doc, r".\.venv\Scripts\python.exe src\run_locality_model_and_review_refresh.py")
    doc.add_paragraph("This orchestration command runs the locality preparation/model/review sequence. It creates processed files. It does not turn a historical result into a live official alert.")

    doc.add_heading("5. Data scripts: learn input → output", 1)
    table(doc, ["File", "Input", "Output / learning point"], [
        ["fetch_multiscale_pollutant_history.py", "OpenAQ fallback or configured source; area/pollutant/date arguments", "Historical multi-pollutant observations with source labels."],
        ["fetch_cpcb_realtime_snapshot.py", "Authorised CPCB/data.gov.in key", "Current official snapshot when access exists; otherwise safely disabled."],
        ["validate_observation_contract.py", "Raw observation table", "Checks canonical fields, pollutants, timestamps and provenance."],
        ["report_multiscale_coverage.py", "Processed history", "Rows, span, gaps, stations and completeness by area/pollutant."],
        ["build_locality_evidence_profile.py", "Coverage and station data", "Explains whether a locality has enough evidence for modelling/review."],
    ], [2.1, 2.0, 2.4])
    callout(doc, "Why six pollutants?", "PM2.5 is the main early forecast target in this prototype, while PM10, NO₂, O₃, SO₂ and CO provide broader air-quality context. Do not claim that all six are independently forecast or that a single number is an official AQI.", PALE_AMBER)

    doc.add_heading("6. Feature engineering: the line of thinking", 1)
    doc.add_paragraph("A model cannot predict the next hour from future information. Feature engineering creates columns that were available at the prediction time.")
    table(doc, ["Feature", "Example", "Reason"], [
        ["Lag", "pm25_lag_1h = PM2.5 measured one hour earlier", "Air quality has time continuity."],
        ["Rolling mean", "mean of previous 3 or 6 hours", "Reduces the effect of one noisy observation."],
        ["Time features", "hour and day_of_week", "Captures patterns like rush-hour or weekly changes."],
        ["Weather", "temperature, humidity, wind, rainfall, pressure", "Weather can influence pollutant dispersion and persistence."],
        ["Target", "pm25_next_1h / next_3h / next_4h", "The future value used only for supervised training/evaluation."],
    ], [1.35, 2.7, 2.45])
    doc.add_paragraph("Read src/features/build_locality_multi_horizon_features.py. The most important safety rule is that all lag/rolling features look backwards in time. The file excludes rows with missing requirements rather than silently inventing values.")

    doc.add_heading("7. ML in simple language", 1)
    table(doc, ["Term", "Meaning in AirSentinel"], [
        ["Persistence baseline", "A simple comparison: predict the next value will be the current value. Any conventional model must beat it to be useful."],
        ["Conventional model", "A non-deep-learning baseline model using tabular lag/weather features. The project deliberately does not promote deep or federated learning."],
        ["Chronological split", "Earlier data is training; later data is testing. This imitates real forecasting and prevents future leakage."],
        ["MAE", "Average absolute forecast error. Lower is better and easy to explain in µg/m³."],
        ["RMSE", "Error metric that penalises larger mistakes more strongly. Lower is better."],
        ["Empirical interval", "A historically derived uncertainty range around the prediction. It is not a guarantee."],
        ["Promotion gate", "A quality and comparison check. A model does not become operational only because it produces a number."],
    ], [1.8, 4.7])
    code_block(doc, r".\.venv\Scripts\python.exe src\models\train_locality_multi_horizon.py")
    code_block(doc, r".\.venv\Scripts\python.exe src\models\predict_locality_multi_horizon.py")
    callout(doc, "What to say to judges", "‘We use a conservative conventional time-series baseline with chronological hold-out evaluation. We show uncertainty and keep new localities gated until they have enough aligned data. We do not claim deep learning or federated learning.’", PALE_GREEN)

    doc.add_page_break()
    doc.add_heading("8. Sudden spike: why the model does not ‘fail silently’", 1)
    doc.add_paragraph("Suppose a model expects PM2.5 around 28 but a later verified value is 38. The difference is called a residual. The anomaly pipeline measures whether a residual is unusually large compared with historical residuals. If it is, it creates a review case.")
    table(doc, ["The code does", "The code deliberately does not do"], [
        ["Keeps a residual and robust residual score.", "Claim that stubble burning, an industry or a person caused the spike."],
        ["Requests corroboration: fresh stations, weather and moderated community context.", "Automatically issue enforcement, public health or source-attribution instructions."],
        ["Adds the case to a queue with an owner/status/action fields.", "Treat a photo, purifier or citizen report as a calibrated outdoor monitor."],
    ], [3.25, 3.25])
    code_block(doc, r".\.venv\Scripts\python.exe src\scoring\build_locality_anomaly_cases.py")
    code_block(doc, r".\.venv\Scripts\python.exe src\authority\build_authority_case_queue.py")

    doc.add_heading("9. Dashboards, API and safety boundary", 1)
    table(doc, ["Surface", "Who uses it", "What it is allowed to do"], [
        ["app.py", "Public/judges", "Show aggregate reviewed information and explanation of uncertainty."],
        ["citizen_app.py", "A consented public user", "Create coarse evidence for pending moderation; no claim of measurement truth."],
        ["authority_app.py", "Authorised reviewer", "Review queue and audited action workflow; production is deliberately view-only."],
        ["src/api/main.py", "Protected service-to-service/client calls", "Validate identity/schema, make controlled durable actions, audit them."],
    ], [1.6, 1.4, 3.5])
    doc.add_paragraph("FastAPI is the backend framework. Streamlit is the dashboard framework. A browser talks to Streamlit for screens; protected services talk to FastAPI for controlled actions. API schemas prevent malformed payloads; security checks refuse an unauthorised role.")

    doc.add_heading("10. Google integrations: understand the adapter pattern", 1)
    doc.add_paragraph("The src/cloud and src/services folders use guarded adapters. That means a feature is disabled unless the required environment configuration exists. This avoids a dashboard pretending a provider call succeeded when the project/key has not been configured.")
    table(doc, ["Adapter", "Real job", "Safe rule"], [
        ["Gemini explainer/photo", "Summarise bounded evidence or describe consented photo context.", "No AQI, cause, enforcement, health or liability decision."],
        ["Firebase store", "Hold pending reviewed citizen evidence.", "Verified identity, deny-by-default rules, retention and moderation."],
        ["BigQuery loader", "Store approved analytical/audit tables.", "Least-privilege roles; no raw personal reports."],
        ["Maps / Earth Engine", "Show locality/satellite context.", "Satellite column/context layer is not surface PM2.5 proof."],
        ["Language services", "Translate or transcribe consented reports.", "Keep consent, language QA and human review."],
        ["Vertex readiness", "Possible managed hosting path after evaluation.", "Not a claim that a superior deep model is deployed."],
    ], [1.75, 2.35, 2.4])

    doc.add_page_break()
    doc.add_heading("11. How to test before saying it works", 1)
    code_block(doc, r".\.venv\Scripts\python.exe -m unittest discover -s tests -v")
    code_block(doc, r".\.venv\Scripts\python.exe -m compileall -q src app.py authority_app.py citizen_app.py")
    doc.add_paragraph("The tests cover API access, audit safety, cloud integration safety, source/promotion gates and long-history ingestion. A test passing is evidence that a defined expected behaviour remains intact; it is not proof that a real government deployment is approved.")
    table(doc, ["Common issue", "Meaning", "First response"], [
        ["‘Path not found’", "Command Prompt is not inside the project folder.", "Run cd C:\\Users\\Rupesh\\AirSentinel\\AirSentinel first."],
        ["‘not recognised’", "A Python script was typed without python.exe or from the wrong folder.", "Use .\\.venv\\Scripts\\python.exe followed by the script path."],
        ["Dashboard port busy", "Another Streamlit process owns 8501/8502/8503.", "Stop it with Ctrl+C or use a different --server.port."],
        ["Disabled Google adapter", "No approved project/key/feature flag was supplied.", "This is safe. Configure through Secret Manager/GCP, not by hard-coding a key."],
    ], [1.55, 2.6, 2.35])

    doc.add_heading("12. Ten questions you should be ready to answer", 1)
    qa = [
        ("What problem is solved?", "A city-to-locality evidence, forecast and review gap—not an attempt to replace the regulator."),
        ("Why not city average only?", "A city average can hide a local industrial edge, village boundary or corridor event."),
        ("Why 1h/3h/4h?", "They give short operational lead time while keeping the research scope conservative."),
        ("Why not use deep learning?", "Data quality, locality coverage and explainability are more important than a complex model claim. The team does not promote unvalidated deep learning."),
        ("How handle a surprise spike?", "Compare later observation to prediction, create a residual case, request corroboration and require human review."),
        ("Can you identify the source?", "No. Evidence can support triage but source attribution requires official investigation."),
        ("What does Gemini do?", "Guarded language/photo context assistance only; it cannot issue AQI, cause or enforcement decisions."),
        ("What is your data fallback?", "OpenAQ development fallback, visibly labelled; official CPCB/data.gov.in path is preferred when authorised."),
        ("How protect users?", "Coarse consented reports, verified identity in production, deny-by-default rules, moderation, retention and audit logs."),
        ("What must happen before pilot use?", "Official/fresh data, station/locality validation, identity and persistence tests, human workflow approval and cloud governance."),
    ]
    for question, answer in qa:
        p = doc.add_paragraph()
        p.add_run(question + " ").bold = True
        p.add_run(answer)

    doc.add_heading("13. Your practical study checklist", 1)
    numbered(doc, [
        "Run the public dashboard and explain every section in your own words.",
        "Open one CSV in data/processed and trace the producing script.",
        "Read src/features/build_locality_multi_horizon_features.py and identify two lags and one weather feature.",
        "Read src/models/train_locality_multi_horizon.py and find the chronological train/test split and metrics.",
        "Read src/scoring/build_locality_anomaly_cases.py and explain why cause_status stays ‘unverified’.",
        "Read src/api/security.py and explain why protected endpoints need a role.",
        "Run tests and say what passing tests mean—and what they do not prove.",
    ])
    doc.core_properties.title = "AirSentinel Code Learning Handbook"
    doc.core_properties.author = "AirSentinel Team"
    doc.save(OUT_LEARN)


def system_guide():
    doc = Document()
    setup(doc, "System working explained")
    cover(doc, "How AirSentinel Works", "The complete working of a locality-aware clean-air review system", "This is the system explanation you can give to teammates, judges and reviewers. It shows the real workflow, the safety boundary and the difference between a working prototype and a live government service.", ["The problem and the users", "End-to-end workflow", "Locality, forecasting and anomaly handling", "Citizen evidence and authority review", "Google Cloud architecture, limits and decision rules"])

    doc.add_heading("1. The system in one sentence", 1)
    callout(doc, "AirSentinel", "AirSentinel helps a verified human reviewer move from city-scale visibility to a specific locality evidence gap or unexpected pollution event, using station data, weather, satellite context and privacy-minimised community evidence—while preserving uncertainty and preventing automatic blame or enforcement.", PALE_TEAL)

    doc.add_heading("2. Problem statement", 1)
    doc.add_paragraph("Major BRICS cities may monitor macro air quality but miss hyper-local and cross-border events such as industrial emissions, agricultural burning or trans-boundary smog. A city average cannot describe every locality, village edge, industrial cluster or transport corridor. The result is a scale gap, a coverage gap, a forecast-failure gap and a coordination gap.")
    table(doc, ["Gap", "What can go wrong", "AirSentinel response"], [
        ["Scale", "Delhi-wide information hides a potential issue near Najafgarh, Rohini, Mehrauli or Safdarjung.", "Uses locality pilot anchors, evidence profiles and nearby-station proxy language."],
        ["Coverage", "No fresh, trusted station evidence supports a precise locality claim.", "Shows evidence gap and asks for coverage review instead of inventing AQI."],
        ["Forecast failure", "A sudden dust/burning/traffic/sensor event differs from expected conditions.", "Creates a cause-unverified residual review case with corroboration steps."],
        ["Coordination", "Stations, weather, satellite and citizen reports are difficult to review together.", "Creates a provenance-labelled, human-reviewed authority queue."],
    ], [1.15, 2.75, 2.55])

    doc.add_heading("3. The people and systems involved", 1)
    table(doc, ["Actor", "What they do", "What they cannot do"], [
        ["Public user", "Views reviewed aggregate conditions and system transparency.", "See raw citizen evidence or take authority action."],
        ["Citizen reporter", "Submits consented coarse text/photo/voice evidence.", "Create an official measurement or accuse a pollution source."],
        ["Authority reviewer", "Examines cases, requests verification and records a human decision.", "Rely on a model alone for enforcement or cause attribution."],
        ["Data/ML team", "Maintains quality gates, source labels, forecasts and tests.", "Label fallback development data as official."],
        ["Google services", "Provide optional storage, AI, maps, satellite or language capability.", "Bypass consent, review, authentication or quality gates."],
    ], [1.45, 2.75, 2.25])

    doc.add_heading("4. End-to-end working flow", 1)
    add_image(doc, ASSETS / "final_end_to_end_workflow.png", "End-to-end AirSentinel workflow: evidence enters through quality gates, then supports forecast, review and auditable action.")
    simple_flow(doc, "What happens in practice", [
        ("Collect", "Download station observations for PM2.5, PM10, NO₂, O₃, SO₂ and CO; collect weather; optionally prepare satellite and moderated community context."),
        ("Label and validate", "Every record keeps source/provenance, time, units, station details and quality/coverage checks."),
        ("Resolve locality", "Relate station evidence to named pilot anchors. Until geometry and physical mapping are approved, it is a nearby-station proxy."),
        ("Forecast", "Use past PM2.5, weather and time features to build research candidates for the next 1, 3 or 4 hours."),
        ("Check mismatch", "When later data arrives, compare it with the prior forecast and assess whether the residual is unusually large."),
        ("Create review case", "Either an evidence gap or an unexpected residual goes into a human authority queue."),
        ("Decide and audit", "A reviewer verifies evidence, chooses a documented next step and records an audit event."),
    ])

    doc.add_heading("5. Data trust: from source to safe use", 1)
    add_image(doc, ASSETS / "final_locality_evidence_gate.png", "Locality evidence gate: the system stops a locality-level claim when coverage, freshness or station relationship evidence is insufficient.")
    table(doc, ["Data type", "Role", "Limit"], [
        ["CPCB/data.gov.in or approved authority export", "Preferred official current/production data source.", "Requires authorised access, station metadata and policy approval."],
        ["OpenAQ", "Historical development fallback used for prototype research.", "Must remain visibly labelled; it is not an authorised CPCB archive."],
        ["Open-Meteo / approved weather", "Weather features/context for dispersion and persistence.", "Weather does not prove a pollution cause."],
        ["Earth Engine/Sentinel-5P", "Satellite context such as broad NO₂/CO/SO₂ signals.", "Vertical-column products are not surface PM2.5 or locality proof."],
        ["Citizen/photo/purifier evidence", "Supporting context for triage and corroboration.", "Not a calibrated outdoor monitoring replacement."],
    ], [1.75, 2.55, 2.15])
    callout(doc, "Locality truth rule", "If there are no fresh, approved or sufficiently related observations, AirSentinel displays a coverage gap and places a review request. It never manufactures a locality AQI to fill the gap.", PALE_AMBER)

    doc.add_heading("6. Forecasting and early warning", 1)
    doc.add_paragraph("The prototype predicts PM2.5 one, three and four hours ahead because that can give a reviewer time to verify conditions and coordinate a response. It uses conventional, explainable tabular models with past PM2.5 lags, rolling averages, time of day and weather features. Earlier data is used for training and later data for testing.")
    table(doc, ["Step", "Meaning"], [
        ["Create lag/weather features", "Only use values available before the prediction time."],
        ["Compare with persistence", "A simple ‘next value equals current value’ baseline ensures the model earns its complexity."],
        ["Chronological test", "Test on later time, not random rows, to imitate a real forecast."],
        ["Show uncertainty", "Display empirical lower/upper ranges based on historical errors."],
        ["Apply promotion gate", "Keep insufficient or weak locality candidates out of operational use."],
    ], [2.1, 4.4])
    callout(doc, "Current honest status", "The locality 1h/3h/4h outputs are historical prototype forecasts, not live official alerts. Najafgarh and Safdarjung need 240 leakage-safe aligned evaluation rows per horizon before controlled baseline comparison; current counts are below that threshold.", PALE_AMBER)

    doc.add_heading("7. How the system handles a sudden spike", 1)
    add_image(doc, ASSETS / "final_evidence_to_action.png", "Evidence-to-action safety flow: unexpected signals become reviewed cases, not automatic source attribution or enforcement.")
    doc.add_paragraph("Example: the forecast estimates PM2.5 of 28 µg/m³ but a later verified observation is 38 µg/m³. AirSentinel records a residual of +10. It compares this residual with historical residuals. If unusual, the system creates an ‘unexpected spike review’ case.")
    numbered(doc, [
        "The system checks whether fresh nearby station readings agree or conflict.",
        "It adds weather and available satellite/context evidence without converting it into proof.",
        "It may show moderated, consented community evidence as supporting context.",
        "It keeps cause_status = unverified. The system does not claim parali, a factory, traffic or a person caused the spike.",
        "An authority reviewer is assigned a verification action. Any action is human initiated and audited.",
    ])
    callout(doc, "Why this is safer", "A forecast can be wrong for legitimate reasons: an unexpected emission, dust event, sensor issue, missing data or conditions not present in training history. The safe response is evidence review, not automatic blame.", PALE_GREEN)

    doc.add_heading("8. Citizen evidence and accessibility", 1)
    doc.add_paragraph("The citizen interface accepts consented text, photos and voice reports. A user may describe smoke, dust or a burning smell in a coarse area. The system minimises personal data, keeps evidence pending review and does not label the report an official sensor reading.")
    table(doc, ["Capability", "How it helps", "Safeguard"], [
        ["Text report", "Adds a local observation for a reviewer to consider.", "Coarse location, consent, moderation and retention rule."],
        ["Photo + Gemini", "Provides a bounded description of visible context.", "No diagnosis, AQI, source attribution or enforcement advice."],
        ["Voice/STT", "Improves access for users who prefer speaking.", "Consent and transcription handling."],
        ["Translation/TTS", "Helps reporters and reviewers use different languages.", "Language QA; reviewer still makes the decision."],
        ["Purifier/home device", "Can suggest a local indoor context signal.", "Cannot replace calibrated outdoor station evidence; only supporting context."],
    ], [1.45, 2.8, 2.25])

    doc.add_heading("9. Authority workflow and audit", 1)
    doc.add_paragraph("The authority queue combines evidence-gap cases and unexpected-spike cases. Every case has a status, recommended next step, owner and audit record. In the local prototype, this demonstrates the workflow. In a real deployment, production actions must use protected FastAPI routes and durable Firestore/BigQuery persistence.")
    table(doc, ["Case type", "Reason", "Reviewer next step"], [
        ["Evidence gap", "There is not enough verified local evidence to support a locality claim.", "Check station coverage, geometry, fresh data and approved access."],
        ["Unexpected spike review", "Observation differs unusually from historical forecast expectation.", "Corroborate with fresh stations, weather and moderated context; assign human review."],
        ["Citizen report review", "A consented report needs moderation and relevance review.", "Validate, keep coarse/pending and decide whether it supports a case."],
    ], [1.7, 2.7, 2.1])
    callout(doc, "No automatic enforcement", "The system may recommend verification or a sustainable response path. It does not issue fines, command inspections, diagnose health impacts or decide that a source is guilty.", PALE_AMBER)

    doc.add_heading("10. Google Cloud architecture", 1)
    add_image(doc, ASSETS / "final_google_cloud_architecture.png", "Separated Google Cloud architecture: public aggregate dashboard, private authority/API services, controlled data stores and optional bounded AI/context services.")
    table(doc, ["Component", "Operational function"], [
        ["Public Cloud Run dashboard", "A separate image/context with only reviewed aggregate information."],
        ["Private Cloud Run API + authority dashboard", "Protected by identity/role checks and used for reviewed actions."],
        ["Firebase Auth + rules", "Verified citizen/reviewer identity and deny-by-default client rules."],
        ["Firestore / BigQuery", "Pending evidence and approved analytical/audit persistence; do not rely on Cloud Run local disk."],
        ["Secret Manager", "Holds API keys/configuration so source code and command lines do not expose them."],
        ["Gemini / Maps / Earth Engine / language APIs", "Feature-flagged, bounded context functions that still obey the human review boundary."],
    ], [2.05, 4.45])

    doc.add_heading("11. What works now versus what needs the team", 1)
    table(doc, ["Already implemented locally", "Needs external team ownership"], [
        ["Data contract, fallbacks, coverage reports, feature/model/anomaly/case pipelines.", "Authorised CPCB/data.gov.in access, station mapping/calibration and official history."],
        ["Public/citizen/authority dashboard code, FastAPI protection, tests and audit workflow.", "Firebase/GCP identities, retention/moderation policy and real persistence verification."],
        ["Cloud Run/Docker/GCP scaffolding and Google adapters.", "GCP project, billing, Secret Manager values, API enablement, deployment and public URLs."],
        ["Reports, flowcharts, pitch deck, demo script and curated submission package.", "Team GitHub repository, video recording and final form links."],
    ], [3.25, 3.25])

    doc.add_heading("12. Limits that make the claim credible", 1)
    bullets(doc, [
        "It is a historical/local prototype decision-support system, not an official AQI publication service.",
        "Nearby station data is a proxy until locality boundaries and physical station relationships are validated.",
        "Satellite, purifier and citizen evidence are contextual signals, not proof or substitute measurements.",
        "A conventional model is only one evidence layer. It can be wrong, so uncertainty and residual review are shown.",
        "No source attribution or enforcement is automated; a competent human authority is needed for operational action.",
        "Google integrations are not claimed live until the team configures the project, secrets, identities and service quotas.",
    ])

    doc.add_page_break()
    doc.add_heading("13. The judge-facing explanation", 1)
    callout(doc, "30-second explanation", "AirSentinel is a locality-aware clean-air early-warning and human-review platform. It combines provenance-labelled outdoor monitoring, weather, satellite context and consented community evidence to find coverage gaps and unexpected pollution changes, forecast one to four hours ahead, and guide a transparent sustainable response. It does not invent hyper-local AQI, automate blame or replace official regulatory judgement.", PALE_TEAL)
    doc.add_heading("Useful source documents", 2)
    bullets(doc, [
        "README.md — local start instructions and current truth.",
        "HANDOFF.md — team status and remaining external tasks.",
        "SYSTEM_ARCHITECTURE.md and docs/FLOWCHARTS.md — architecture/workflow details.",
        "docs/DATA_SOURCE_AND_FALLBACK_POLICY.md — provenance and approved-source rules.",
        "docs/LOCALITY_MODEL_CARD.md — model limits, metrics and promotion boundary.",
        "deploy/README_GOOGLE_CLOUD.md — team-owned cloud deployment plan.",
    ])
    doc.core_properties.title = "How AirSentinel Works"
    doc.core_properties.author = "AirSentinel Team"
    doc.save(OUT_SYSTEM)


if __name__ == "__main__":
    team_guide()
    learning_guide()
    system_guide()
    print(f"Created {OUT_TEAM}")
    print(f"Created {OUT_LEARN}")
    print(f"Created {OUT_SYSTEM}")
