from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


BRIEF_FILE = Path(
    r"C:\Users\Rupesh\Documents\Codex\2026-08-25\referenced-chatgpt-conversation-this-is-an"
) / "AirSentinel_Team_Workflow_and_Plan.docx"


def set_run_font(run, size=9.5, bold=False, color="000000"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def replace_paragraph(paragraph, text, size=11, bold=False, color="000000"):
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def replace_cell(cell, text, bold=False):
    paragraph = cell.paragraphs[0]
    replace_paragraph(paragraph, text, size=9.5, bold=bold)
    paragraph.paragraph_format.space_after = Pt(0)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


document = Document(BRIEF_FILE)
paragraphs = document.paragraphs

# Replace Phase 2 status and results.
replace_paragraph(paragraphs[29], "Phase 2: Forecasting baseline - COMPLETE", size=13, bold=True, color="2E74B5")
replace_paragraph(
    paragraphs[30],
    "Chennai one-hour PM2.5 forecast is trained, evaluated and shown in the dashboard.",
    size=11,
    color="5B6573",
)
replace_paragraph(
    paragraphs[31],
    "699 model-ready Chennai rows were built from PM2.5 lags, rolling averages and historical weather.",
)
replace_paragraph(
    paragraphs[32],
    "The time-ordered Random Forest achieved MAE 4.067 and RMSE 5.694, improving on the persistence baseline.",
)
replace_paragraph(
    paragraphs[33],
    "The dashboard shows a clearly labelled historical validation: 24.9 µg/m³ predicted versus 25.0 µg/m³ recorded.",
)

# Activate Phase 3.
replace_paragraph(paragraphs[34], "Phase 3: Sustainability intelligence - CURRENT", size=13, bold=True, color="2E74B5")
replace_paragraph(
    paragraphs[35],
    "Turn verified current conditions and forecasts into transparent, responsible sustainability guidance.",
    size=11,
    color="5B6573",
)
replace_paragraph(
    paragraphs[36],
    "Build city risk labels from current PM2.5 severity plus a small weather-persistence adjustment; keep the formula visible.",
)
replace_paragraph(
    paragraphs[37],
    "Show sustainable actions by risk level and keep 'data unavailable' where no fresh PM2.5 reading exists. Do not claim a proven pollution source.",
)

# Update the current-status table.
status_table = document.tables[2]
replace_cell(status_table.rows[5].cells[0], "Forecast model", bold=True)
replace_cell(
    status_table.rows[5].cells[1],
    "Chennai baseline trained, evaluated and displayed as a historical validation forecast",
)
replace_cell(status_table.rows[5].cells[2], "Complete")
shade_cell(status_table.rows[5].cells[2], "E4F4EA")

new_row = status_table.add_row().cells
replace_cell(new_row[0], "Sustainability intelligence", bold=True)
replace_cell(
    new_row[1],
    "Transparent city risk labels and sustainable action guidance",
)
replace_cell(new_row[2], "Current")
shade_cell(new_row[2], "FFF4D6")

# Update the final Phase 3 action block without changing the document layout.
replace_paragraph(paragraphs[56], "Phase 3 first milestone", size=16, bold=True, color="2E74B5")
replace_paragraph(
    paragraphs[58],
    "Create transparent city risk scores from verified current conditions and show why each label was assigned.",
)
replace_paragraph(
    paragraphs[59],
    ".\\.venv\\Scripts\\python.exe src\\scoring\\build_city_risk_scores.py",
    size=9.5,
    color="0B2545",
)
replace_paragraph(
    paragraphs[60],
    "Output: data\\processed\\city_risk_scores.csv",
    size=9.5,
    color="0B2545",
)
replace_paragraph(
    paragraphs[61],
    "Then add the risk table and method explanation to app.py.",
    size=10.5,
)
replace_paragraph(paragraphs[62], "", size=10.5)

replace_paragraph(
    paragraphs[66],
    "Phase 3 scoring script: src\\scoring\\build_city_risk_scores.py",
)
replace_paragraph(
    paragraphs[67],
    "Phase 3 output: data\\processed\\city_risk_scores.csv",
)

document.core_properties.title = "AirSentinel - Team Workflow, Phases and Work Split"
document.core_properties.subject = "Phase 2 completion and Phase 3 sustainability intelligence"
document.save(BRIEF_FILE)
print(BRIEF_FILE)
