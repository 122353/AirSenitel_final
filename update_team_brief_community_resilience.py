"""Append the community-scale resilience decision to the AirSentinel team brief."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


DOCUMENT_PATH = Path(
    r"C:\Users\Rupesh\Documents\Codex\2026-08-25\referenced-chatgpt-conversation-this-is-an"
) / "AirSentinel_Team_Workflow_and_Plan.docx"


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def add_number(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run(text)


document = Document(DOCUMENT_PATH)

document.add_page_break()
document.add_heading("Team decision update: community-scale resilience", level=1)

document.add_heading("Revised main problem", level=2)
document.add_paragraph(
    "AirSentinel is not only a city PM2.5 forecast dashboard. It is a pan-India "
    "decision-support prototype for detecting, explaining and triaging air-pollution "
    "risk from the national level down to districts, small cities, villages, wards, "
    "road corridors and other local hotspots."
)
document.add_paragraph(
    "The challenge is that public outdoor monitoring is uneven and a local event can "
    "change conditions faster than a forecast expects. The prototype therefore combines "
    "forecasting with anomaly detection and privacy-safe community evidence. Its role is "
    "to give authorities an earlier, explainable case for review - not to claim that it "
    "has proven the cause of pollution."
)

document.add_heading("Operating model", level=2)
add_number(document, "India overview: show available state, district and city conditions.")
add_number(document, "Local hotspot layer: drill into a ward, village, corridor, industrial edge or grid cell when data exists.")
add_number(document, "Forecast layer: estimate likely PM2.5 three and four hours ahead where historical coverage is adequate.")
add_number(document, "Unexpected-spike layer: compare new observations with the forecast and recent local pattern.")
add_number(document, "Evidence and response layer: combine verified stations with contextual signals and community reports, then create an authority-review case.")

document.add_heading("Handling a sudden local spike", level=2)
document.add_paragraph(
    "Example: a South Delhi forecast suggests PM2.5 near 28, but a verified new reading "
    "is 37-38. AirSentinel must not hide the miss. It labels the event as an unexpected "
    "observed spike and follows this response flow:"
)
add_number(document, "Calculate the difference between the forecast and the new monitored value, plus deviation from the recent local baseline.")
add_number(document, "Check data freshness and nearby verified stations to rule out a single stale or faulty reading.")
add_number(document, "Look for corroboration from wind/rain context, satellite/aerosol context, opted-in sensors and clustered citizen reports.")
add_number(document, "Show: 'Unexpected local spike - cause unverified', including evidence and confidence.")
add_number(document, "Send the case to an authority/reviewer for inspection or communication. Record any later verified cause and response outcome.")

document.add_paragraph(
    "A spike alone cannot prove parali, waste burning, traffic, construction or an industrial source. Those may be displayed only as possible contributors when multiple independent signals support them."
)

document.add_heading("Community, purifier and citizen data", level=2)
document.add_paragraph(
    "AirSentinel can accept opt-in user evidence, but it must keep evidence types separate and privacy-safe.")
add_bullet(document, "Verified outdoor station data is the strongest evidence for public outdoor conditions.")
add_bullet(document, "Weather and satellite data provide context; they do not prove a local source by themselves.")
add_bullet(document, "Citizen voice, text and photo reports are useful evidence, not verified PM2.5 measurements.")
add_bullet(document, "Connected purifier or home-sensor data may be used only with explicit consent, supported vendor/account access and clear indoor-data labelling.")
add_bullet(document, "Store an approximate ward/grid location instead of a raw home address; remove account tokens from analytics; provide deletion and withdrawal of consent.")

document.add_paragraph(
    "There is no single API that exposes live readings from every purifier brand. A later prototype connector can support selected opt-in ecosystems. These readings are usually indoor measurements, so they cannot replace outdoor AQI or be averaged directly into a city PM2.5 value without quality checks, calibration and aggregation."
)

document.add_heading("Revised next phases", level=2)
add_bullet(document, "Phase 2.5: direct 3-hour and 4-hour forecasts; compare time-aware model baselines and communicate uncertainty.")
add_bullet(document, "Phase 3.5: unexpected-spike detector using forecast residuals and local baseline deviation; test it on held-out historical periods.")
add_bullet(document, "Phase 4: BigQuery and Cloud Run, with a locality/ward/grid evidence schema and confidence fields.")
add_bullet(document, "Phase 5: Gemini explanations and Google Earth Engine context, never used as a numeric predictor or source proof.")
add_bullet(document, "Phase 6: Firebase citizen reports, photo evidence and optional authorised home-sensor/purifier connectors.")
add_bullet(document, "Phase 7: Hindi/English voice, translation, text-to-speech and the authority/citizen assistant.")

document.add_heading("Integrity commitments", level=2)
document.add_paragraph(
    "Where a locality has no verified outdoor monitor, AirSentinel will show limited evidence or lower confidence instead of inventing AQI. The prototype will describe its output as decision support, not a guaranteed prediction, proven source attribution or substitute for official government measurement."
)

document.add_heading("Government-facing pilot requirements", level=1)
document.add_paragraph(
    "AirSentinel should complement CPCB, SPCB/PCC, ULB, district and NCAP workflows. It is not a replacement for official AQI, statutory monitoring, enforcement authority or laboratory verification. The team must not make unverified claims about cybersecurity weaknesses in government systems and must never attempt to access, scan or bypass them."
)
document.add_heading("What AirSentinel adds", level=2)
add_bullet(document, "Local evidence retained alongside city summaries: station, hotspot grid and locality confidence.")
add_bullet(document, "An operational anomaly queue: case ID, evidence, confidence, suggested next step, assigned owner, due time and outcome.")
add_bullet(document, "A clear separation between the public view and an authority view with role-based access.")
add_bullet(document, "Auditable model/rule versions, data freshness and uncertainty rather than unexplained scores.")

document.add_heading("Government operator case flow", level=2)
add_number(document, "A verified measurement, anomaly or corroborated report creates a case with an evidence snapshot.")
add_number(document, "The system displays source tier, data freshness, confidence and recommended review step.")
add_number(document, "A human operator assigns the case, records the action and closes it with an outcome.")
add_number(document, "The team evaluates forecast accuracy, anomaly false alerts, coverage and time to triage from completed cases.")

document.add_heading("Safeguards the pilot must demonstrate", level=2)
add_bullet(document, "A city average must not hide station/local hotspot variation; show location resolution and confidence.")
add_bullet(document, "Sparse coverage must remain visible; never create an official AQI where evidence is missing.")
add_bullet(document, "A forecast miss must create an unexpected-spike case, not be silently ignored.")
add_bullet(document, "One anonymous report cannot trigger enforcement or public source attribution; use rate limits, moderation and corroboration.")
add_bullet(document, "Indoor purifier data must be labelled separately from outdoor monitoring and accepted only with consent.")
add_bullet(document, "No private address, account token or raw device ID may appear in analytics or the public dashboard.")
add_bullet(document, "Automated suggestions require human review before enforcement or public attribution.")

for paragraph in document.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Aptos"
        if run.font.size is None:
            run.font.size = Pt(10.5)

document.save(DOCUMENT_PATH)
print(f"Updated: {DOCUMENT_PATH}")
