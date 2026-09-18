"""Append the multi-scale, multi-pollutant architecture to the team brief once."""

from pathlib import Path

from docx import Document


DOCUMENT_PATH = Path(
    r"C:\Users\Rupesh\Documents\Codex\2026-08-25\referenced-chatgpt-conversation-this-is-an"
) / "AirSentinel_Team_Workflow_and_Plan.docx"
HEADING = "Multi-scale and multi-pollutant architecture"


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


document = Document(DOCUMENT_PATH)
if any(paragraph.text == HEADING for paragraph in document.paragraphs):
    print("Multi-scale section already exists; no change made.")
else:
    document.add_page_break()
    document.add_heading(HEADING, level=1)
    document.add_paragraph(
        "Chennai remains the first completed data-and-model proof, not the product scope. "
        "The scalable architecture works from India and state/city views down to a "
        "district, small city, village, ward, road corridor or hotspot grid when evidence exists. "
        "Every value must display its source, timestamp, locality resolution, freshness and confidence."
    )

    document.add_heading("Priority pollutant inputs", level=2)
    add_bullet(document, "Core: PM2.5, PM10, nitrogen dioxide (NO2), sulphur dioxide (SO2), carbon monoxide (CO) and ozone (O3).")
    add_bullet(document, "Optional where coverage is reliable: ammonia (NH3) and benzene (C6H6).")
    add_bullet(document, "SO2 is the standard ambient pollutant. SO3 is not substituted as a standard AirSentinel target.")
    add_bullet(document, "Weather context: temperature, humidity, wind, rainfall, pressure and official nowcast/warning information where access permits.")

    document.add_heading("Source and evidence policy", level=2)
    add_bullet(document, "Primary: authorised CPCB/data.gov.in station records and official IMD weather/nowcast where access is authorised.")
    add_bullet(document, "Development connector: OpenAQ, always retaining provider/provenance and never automatically described as a government source.")
    add_bullet(document, "Context: satellite/aerosol features and weather patterns, never source proof on their own.")
    add_bullet(document, "Community input: opt-in device, voice, photo and manual reports; separate from verified outdoor station data.")

    document.add_heading("Deep-learning decision", level=2)
    document.add_paragraph(
        "The future advanced model is a multi-task spatiotemporal model. It will learn "
        "from station/locality sequences, neighbouring locations, weather and satellite context, "
        "with separate pollutant targets and forecast horizons. It will not be trained simply "
        "because it sounds advanced."
    )
    add_bullet(document, "Data gate: at least 6-12 months across seasons, multiple stations per pilot region, aligned timestamps and quality/missingness flags.")
    add_bullet(document, "Evaluation gate: time-ordered cross-validation plus a final untouched test period for every location and pollutant.")
    add_bullet(document, "Acceptance gate: the deep model must improve on persistence and gradient-boosting baselines without worsening low-coverage localities.")

    document.add_heading("Immediate implementation", level=2)
    add_bullet(document, "Run multi-pollutant station discovery for Delhi NCR, Mehrauli, Safdarjung Enclave, Kochi and current macro-city pilots.")
    add_bullet(document, "Download the available hourly pollutant history and create a coverage matrix before fitting a pan-India model.")
    add_bullet(document, "Add direct 3-hour and 4-hour forecasts, then anomaly detection and a professional authority dashboard.")

    document.save(DOCUMENT_PATH)
    print(f"Updated: {DOCUMENT_PATH}")
